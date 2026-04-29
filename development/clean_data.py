import pandas as pd
import re
import os
import glob


def get_latest_file(data_dir, pattern_xlsx, pattern_csv=None):
    """Helper to find the most recently created/named file matching a pattern."""
    files = glob.glob(os.path.join(data_dir, pattern_xlsx))
    if pattern_csv:
        files.extend(glob.glob(os.path.join(data_dir, pattern_csv)))
    
    # Filter out enriched ones just in case we are matching `*-wiki.csv`
    files = [f for f in files if not f.endswith('-wiki-enriched.csv')]
    
    if not files:
        raise FileNotFoundError(f"No files found matching {pattern_xlsx} or {pattern_csv}")
    files.sort(reverse=True)
    return files[0]


def rewrite_collection_rows(df: pd.DataFrame, title_col: str, notes_col: str) -> pd.DataFrame:
    """Replace collection headers with the actual per-game titles from notes."""
    if title_col not in df.columns or notes_col not in df.columns:
        return df

    title_values = df[title_col].astype("string").str.strip()
    notes_values = df[notes_col].astype("string").str.strip()
    has_title = title_values.notna() & title_values.ne("")
    has_notes = notes_values.notna() & notes_values.ne("")

    collection_names = title_values.where(has_title).ffill()
    is_continuation = ~has_title & has_notes
    next_is_continuation = is_continuation.shift(-1).fillna(False)
    is_start = has_title & has_notes & next_is_continuation
    in_collection = is_start | is_continuation

    new_titles = title_values.copy()
    new_notes = notes_values.copy()

    new_titles.loc[in_collection] = notes_values.loc[in_collection]
    new_notes.loc[in_collection] = "Part of collection: " + collection_names.loc[in_collection]

    isolated_missing_title = ~in_collection & ~has_title & has_notes
    new_titles.loc[isolated_missing_title] = notes_values.loc[isolated_missing_title]

    df[title_col] = new_titles
    df[notes_col] = new_notes
    return df


def drop_wiki_bundle_headers(df: pd.DataFrame) -> pd.DataFrame:
    """Drop stale bundle header rows when the individual games are already present."""
    required_cols = {'Title', 'Date', 'Source Links'}
    if df.empty or not required_cols.issubset(df.columns):
        return df

    df = df.copy()
    df['Title'] = df['Title'].astype("string").str.strip()
    df['is_bundle_header'] = df['Title'].str.contains(
        r'\b(?:collection|trilogy)\b',
        case=False,
        regex=True,
        na=False,
    )

    group_cols = ['Date', 'Source Links']
    group_sizes = df.groupby(group_cols, dropna=False)['Title'].transform('size')
    header_counts = df.groupby(group_cols, dropna=False)['is_bundle_header'].transform('sum')
    keep_mask = ~(df['is_bundle_header'] & (group_sizes > 1) & (header_counts > 0))

    return df.loc[keep_mask].drop(columns='is_bundle_header').reset_index(drop=True)


def clean_gsheets():
    """
    Cleans the Google Sheets data.

    Data Key for the 'TYPE' column (extracted from original headers):
      - pack: packs, in-game content, DLC
      - dupe: dupe games, packs, in-game content, DLC
      - free: free games
      - other: other unique games (another days, different number of days, DLC's, fails)
      - pack2: other packs, in-game content, DLC
      - dupe2: other dupe games, packs, in-game content, DLC
      - mobile: mobile unique games
      - pack3: mobile packs, in-game content, DLC
      - dupe3: mobile dupe games, packs, in-game content, DLC
      - next: next gifts (future/upcoming games)
    """
    # Find the latest Google Sheets file dynamically
    data_dir = os.path.join(os.path.dirname(__file__), '../data')
    file_path = get_latest_file(data_dir, "20*-gsheets.xlsx", "20*-gsheets.csv")
    
    if file_path.endswith('.xlsx'):
        df = pd.read_excel(file_path, engine='openpyxl', skiprows=15)

        import openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb.active

        colors = []
        for row_idx in range(17, 17 + len(df)):
            cell = sheet.cell(row=row_idx, column=6) # 6th column is 'NAME'
            rgb = cell.fill.start_color.rgb if cell.fill and cell.fill.start_color else None
            colors.append(rgb if rgb and isinstance(rgb, str) else 'unknown')
    else:
        df = pd.read_csv(file_path, skiprows=15)
        colors = ['unknown'] * len(df)

    # Keep the first 7 columns and rename to standard names
    df = df.iloc[:, 0:7].copy()
    df.columns = ['FROM', 'TO', 'DAY', 'DAYS', 'TYPE', 'Title', 'NOTES']

    df['COLOR_CATEGORY'] = colors

    # Slice from the first "next" and exclude trailing rows
    first_next_idx = df.index[df['TYPE'] == 'next'].tolist()
    if first_next_idx:
        start_idx = first_next_idx[0]
        df = df.loc[start_idx:].copy()

    # Drop the last row if it repeats the header
    if len(df) > 0 and df.iloc[-1]['FROM'] == 'FROM':
        df = df.iloc[:-1].copy()

    # rewrite collection rows before forward-filling so the real game title is kept
    df = rewrite_collection_rows(df, 'Title', 'NOTES')

    # Forward-fill event-related columns for days with multiple games (Do not ffill TYPE)
    cols_to_ffill = ['FROM', 'TO', 'DAY', 'DAYS']
    for col in cols_to_ffill:
        df[col] = df[col].ffill()

    # Clean string values, remove empties, and filter out future placeholders
    df = df[df['Title'].notna()].copy()
    df['Title'] = df['Title'].astype("string").str.strip()
    df = df[~df['Title'].isin(['', '-', 'nan'])]
    df = df[df['TYPE'] != '*']

    # We reset index for cleaner look
    df = df.reset_index(drop=True)

    return df

def clean_wiki():
    """
    Loads and cleans the scraped Wikipedia dataset.

    Returns:
        pd.DataFrame: A DataFrame with exploded titles replacing bundle text.
    """
    # Load the Wikipedia CSV file dynamically
    data_dir = os.path.join(os.path.dirname(__file__), '../data')
    file_path = get_latest_file(data_dir, "20*-wiki.csv")
    df = pd.read_csv(file_path)

    # Explode rows where Title has multiple games separated by \n
    df['Title'] = df['Title'].astype("string").str.split('\n')
    df = df.explode('Title')
    df['Title'] = df['Title'].astype("string").str.strip()
    df = drop_wiki_bundle_headers(df)

    return df

def normalize_title_series(s):
    """Normalize a title Series for robust merging: lowercase, remove special characters, and strip extra spaces."""
    return s.astype("string").str.lower().str.replace(r'[^a-z0-9\s]', '', regex=True).str.replace(r'\s+', ' ', regex=True).str.strip()

def main():
    """
    Main function to execute the data cleaning pipeline.

    Loads and cleans Google Sheets and Wikipedia datasets separately, normalizes
    titles to perform a left join merging Wikipedia metadata into the primary
    Google Sheets dataset, and saves the result to a CSV file.
    """
    print("Cleaning Google Sheets data...")
    gsheets_df = clean_gsheets()

    print("Cleaning Wikipedia data...")
    wiki_df = clean_wiki()

    # Create temporary normalized title columns for merging
    gsheets_df['merge_title'] = normalize_title_series(gsheets_df['Title'])
    wiki_df['merge_title'] = normalize_title_series(wiki_df['Title'])

    print("Merging datasets...")
    # LEFT JOIN: keep all from gsheets (primary source), enrich with wiki info where possible
    merged_df = pd.merge(gsheets_df, wiki_df, on='merge_title', how='left', suffixes=('_gsheets', '_wiki'))

    # Drop the temporary merge column
    merged_df = merged_df.drop('merge_title', axis=1)

    print(f"Merge complete. Result shape: {merged_df.shape}")
    print(f"Games successfully matched with Wikipedia data: {merged_df['Title_wiki'].notna().sum()}")

    # Save output to data directory
    output_path = os.path.join(os.path.dirname(__file__), '../data/cleaned_merged_data.csv')
    merged_df.to_csv(output_path, index=False)
    print(f"Saved merged dataset to: {output_path}")

if __name__ == "__main__":
    main()
